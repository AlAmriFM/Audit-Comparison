#!/usr/bin/env python3
"""
ISMS Admin Account Audit Comparator

Standalone offline desktop tool for comparing two quarterly admin account
Excel submissions and generating a self-contained HTML audit report.
"""

from __future__ import annotations

import datetime as _dt
import html
import json
import os
import re
import shutil
import sys
import traceback
import webbrowser
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - shown in GUI at runtime
    load_workbook = None

try:
    import xlrd
except ImportError:  # pragma: no cover - optional dependency
    xlrd = None

try:
    from PIL import Image, ImageEnhance, ImageOps
except ImportError:  # pragma: no cover - optional dependency
    Image = None
    ImageEnhance = None
    ImageOps = None

try:
    import pytesseract
except ImportError:  # pragma: no cover - optional dependency
    pytesseract = None


ROLE_LABELS = {
    "user_id": "User ID",
    "rights_type": "Type of admin rights",
    "granted_by": "Granted by",
    "purpose": "Purpose",
    "shared_single": "Shared / single-user",
    "expiration": "Expiration",
    "request_id": "Request ID",
}

COMPARABLE_ROLES = [
    "rights_type",
    "granted_by",
    "purpose",
    "shared_single",
    "expiration",
    "request_id",
]

SCREENSHOT_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}

ROLE_ALIASES = {
    "user_id": [
        "user id",
        "userid",
        "user",
        "username",
        "user name",
        "account",
        "account id",
        "admin account",
        "login id",
        "login",
        "id",
    ],
    "rights_type": [
        "type of admin rights",
        "admin rights",
        "rights",
        "role",
        "privilege",
        "privileges",
        "access type",
        "admin type",
        "permission",
    ],
    "granted_by": [
        "granted by",
        "approved by",
        "access granted by",
        "grantor",
        "approver",
        "authorized by",
        "authorised by",
    ],
    "purpose": [
        "purpose",
        "business purpose",
        "reason",
        "justification",
        "access purpose",
        "description",
    ],
    "shared_single": [
        "shared",
        "shared account",
        "single user",
        "single-user",
        "individual",
        "named account",
        "account type",
    ],
    "expiration": [
        "expiration",
        "expiry",
        "expiry date",
        "expiration date",
        "user expiration",
        "valid until",
        "end date",
        "validity",
    ],
    "request_id": [
        "request id",
        "requestid",
        "request",
        "change request",
        "ticket",
        "ticket number",
        "cr",
        "cr number",
        "change id",
        "reference",
    ],
}


@dataclass
class SheetTable:
    name: str
    headers: List[str]
    rows: List[Dict[str, Any]]
    header_row: int
    roles: Dict[str, Optional[str]] = field(default_factory=dict)
    skipped: bool = False


@dataclass
class UserRecord:
    key: str
    display_id: str
    values: Dict[str, str]
    raw: Dict[str, Any]


@dataclass
class ScreenshotUserCheck:
    user_id: str
    present: bool
    evidence: str = ""
    match_type: str = ""


@dataclass
class ScreenshotCheckResult:
    system_name: str
    screenshot_path: Optional[str] = None
    status: str = "not_run"
    message: str = ""
    users: List[ScreenshotUserCheck] = field(default_factory=list)
    ocr_preview: str = ""

    @property
    def missing_count(self) -> int:
        return sum(1 for user in self.users if not user.present)

    @property
    def checked_count(self) -> int:
        return len(self.users)


@dataclass
class SystemResult:
    name: str
    additions: List[UserRecord] = field(default_factory=list)
    removals: List[UserRecord] = field(default_factory=list)
    changes: List[Dict[str, Any]] = field(default_factory=list)
    clean: bool = False
    skipped: bool = False
    skip_reason: str = ""
    screenshot_check: Optional[ScreenshotCheckResult] = None


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\n", " ").replace("\r", " ")
    text = re.sub(r"[_\-/:]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip().lower()
    return text


def normalize_user_id(value: Any) -> str:
    return re.sub(r"\s+", " ", cell_to_text(value)).strip().lower()


def cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, _dt.date):
        return value.isoformat()
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def has_value(value: Any) -> bool:
    return cell_to_text(value) != ""


def unique_headers(raw_headers: Iterable[Any]) -> List[str]:
    headers: List[str] = []
    seen: Dict[str, int] = {}
    for index, raw in enumerate(raw_headers, start=1):
        base = cell_to_text(raw) or f"Column {index}"
        base = re.sub(r"\s+", " ", base).strip()
        count = seen.get(base.lower(), 0)
        seen[base.lower()] = count + 1
        headers.append(base if count == 0 else f"{base} ({count + 1})")
    return headers


def score_header_row(row: Tuple[Any, ...]) -> int:
    score = 0
    populated = 0
    for value in row:
        text = cell_to_text(value)
        if not text:
            continue
        populated += 1
        norm = normalize_header(text)
        if len(text) <= 80:
            score += 2
        if any(alias == norm or alias in norm for aliases in ROLE_ALIASES.values() for alias in aliases):
            score += 5
    return score + populated


def detect_header_row(rows: List[Tuple[Any, ...]]) -> int:
    candidates = rows[:8] or [()]
    best_index = 0
    best_score = -1
    for index, row in enumerate(candidates):
        score = score_header_row(row)
        if score > best_score:
            best_index = index
            best_score = score
    return best_index


def trim_empty_tail(row: Tuple[Any, ...], width: int) -> List[Any]:
    values = list(row[:width])
    if len(values) < width:
        values.extend([""] * (width - len(values)))
    return values


def detect_roles(headers: List[str]) -> Dict[str, Optional[str]]:
    roles: Dict[str, Optional[str]] = {role: None for role in ROLE_LABELS}
    used: set[str] = set()
    normalized = [(header, normalize_header(header)) for header in headers]
    for role, aliases in ROLE_ALIASES.items():
        best_header = None
        best_score = 0
        for header, norm in normalized:
            if header in used:
                continue
            score = 0
            for alias in aliases:
                if norm == alias:
                    score = max(score, 100)
                elif norm.replace(" ", "") == alias.replace(" ", ""):
                    score = max(score, 90)
                elif alias in norm:
                    score = max(score, 65 + min(len(alias), 20))
            if score > best_score:
                best_header = header
                best_score = score
        if best_header and best_score >= 60:
            roles[role] = best_header
            used.add(best_header)
    return roles


def load_workbook_tables(path: Path) -> Dict[str, SheetTable]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return load_xlsx_tables(path)
    if suffix == ".xls":
        return load_xls_tables(path)
    raise ValueError(f"Unsupported file type: {path.suffix}. Please select .xlsx or .xls.")


def load_xlsx_tables(path: Path) -> Dict[str, SheetTable]:
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed. Run: pip install -r requirements.txt")
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    tables: Dict[str, SheetTable] = {}
    for sheet in workbook.worksheets:
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        tables[sheet.title] = rows_to_table(sheet.title, rows)
    workbook.close()
    return tables


def load_xls_tables(path: Path) -> Dict[str, SheetTable]:
    if xlrd is None:
        raise RuntimeError("xlrd is not installed. Run: pip install -r requirements.txt")
    workbook = xlrd.open_workbook(str(path))
    tables: Dict[str, SheetTable] = {}
    for sheet in workbook.sheets():
        rows = [tuple(sheet.row_values(row_index)) for row_index in range(sheet.nrows)]
        tables[sheet.name] = rows_to_table(sheet.name, rows)
    return tables


def rows_to_table(name: str, rows: List[Tuple[Any, ...]]) -> SheetTable:
    if not rows:
        return SheetTable(name=name, headers=[], rows=[], header_row=1, roles={role: None for role in ROLE_LABELS})
    header_index = detect_header_row(rows)
    header_row = rows[header_index]
    last_col = len(header_row)
    for row in rows[header_index + 1 :]:
        for idx, value in enumerate(row, start=1):
            if has_value(value):
                last_col = max(last_col, idx)
    headers = unique_headers(trim_empty_tail(header_row, last_col))
    data_rows: List[Dict[str, Any]] = []
    for raw_row in rows[header_index + 1 :]:
        values = trim_empty_tail(raw_row, len(headers))
        if not any(has_value(value) for value in values):
            continue
        data_rows.append(dict(zip(headers, values)))
    return SheetTable(
        name=name,
        headers=headers,
        rows=data_rows,
        header_row=header_index + 1,
        roles=detect_roles(headers),
    )


def build_user_map(table: SheetTable) -> Dict[str, UserRecord]:
    user_col = table.roles.get("user_id")
    if not user_col:
        return {}
    records: Dict[str, UserRecord] = {}
    for row in table.rows:
        key = normalize_user_id(row.get(user_col))
        if not key:
            continue
        display = cell_to_text(row.get(user_col))
        values: Dict[str, str] = {}
        for role, label in ROLE_LABELS.items():
            header = table.roles.get(role)
            values[role] = cell_to_text(row.get(header)) if header else ""
        if key in records:
            records[key].display_id = display or records[key].display_id
        else:
            records[key] = UserRecord(key=key, display_id=display, values=values, raw=row)
    return records


class MappingDialog(tk.Toplevel):
    def __init__(self, parent: tk.Tk, sheet_name: str, previous: SheetTable, current: SheetTable):
        super().__init__(parent)
        self.title(f"Map columns - {sheet_name}")
        self.resizable(True, True)
        self.result: Optional[Tuple[bool, Dict[str, Optional[str]], Dict[str, Optional[str]]]] = None
        self.previous_vars: Dict[str, tk.StringVar] = {}
        self.current_vars: Dict[str, tk.StringVar] = {}

        self.transient(parent)
        self.grab_set()
        self.protocol("WM_DELETE_WINDOW", self.skip)

        outer = ttk.Frame(self, padding=18)
        outer.grid(row=0, column=0, sticky="nsew")
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)
        outer.columnconfigure(0, weight=1)

        title = ttk.Label(outer, text=f"Column mapping needed for: {sheet_name}", style="Title.TLabel")
        title.grid(row=0, column=0, sticky="w")
        help_text = (
            "The User ID column is required. Review the detected mappings below, "
            "choose columns from the dropdowns, or skip this worksheet."
        )
        ttk.Label(outer, text=help_text, wraplength=760).grid(row=1, column=0, sticky="we", pady=(6, 14))

        body = ttk.Frame(outer)
        body.grid(row=2, column=0, sticky="nsew")
        body.columnconfigure(0, weight=1)
        body.columnconfigure(1, weight=1)
        self._mapping_group(body, "Previous quarter", previous, self.previous_vars, 0)
        self._mapping_group(body, "Current quarter", current, self.current_vars, 1)

        buttons = ttk.Frame(outer)
        buttons.grid(row=3, column=0, sticky="e", pady=(18, 0))
        ttk.Button(buttons, text="Skip worksheet", command=self.skip).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(buttons, text="Use mappings", command=self.accept, style="Accent.TButton").grid(row=0, column=1)

        self.bind("<Escape>", lambda _event: self.skip())
        self.bind("<Return>", lambda _event: self.accept())
        self.update_idletasks()
        self.minsize(min(self.winfo_width(), 900), min(self.winfo_height(), 650))

    def _mapping_group(
        self,
        parent: ttk.Frame,
        label: str,
        table: SheetTable,
        vars_by_role: Dict[str, tk.StringVar],
        column: int,
    ) -> None:
        group = ttk.LabelFrame(parent, text=label, padding=12)
        group.grid(row=0, column=column, sticky="nsew", padx=(0, 8) if column == 0 else (8, 0))
        group.columnconfigure(1, weight=1)
        options = ["(none)"] + table.headers
        ttk.Label(group, text=f"Detected header row: {table.header_row}").grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 8)
        )
        for row_index, (role, role_label) in enumerate(ROLE_LABELS.items(), start=1):
            ttk.Label(group, text=role_label).grid(row=row_index, column=0, sticky="w", padx=(0, 8), pady=3)
            initial = table.roles.get(role) or "(none)"
            var = tk.StringVar(value=initial if initial in options else "(none)")
            vars_by_role[role] = var
            combo = ttk.Combobox(group, textvariable=var, values=options, state="readonly", width=30)
            combo.grid(row=row_index, column=1, sticky="we", pady=3)

    def _collect(self, vars_by_role: Dict[str, tk.StringVar]) -> Dict[str, Optional[str]]:
        mapping: Dict[str, Optional[str]] = {}
        for role, var in vars_by_role.items():
            value = var.get()
            mapping[role] = None if value == "(none)" else value
        return mapping

    def accept(self) -> None:
        previous = self._collect(self.previous_vars)
        current = self._collect(self.current_vars)
        if not previous.get("user_id") or not current.get("user_id"):
            messagebox.showwarning(
                "User ID required",
                "Please select a User ID column for both files, or skip this worksheet.",
                parent=self,
            )
            return
        self.result = (False, previous, current)
        self.destroy()

    def skip(self) -> None:
        self.result = (True, {}, {})
        self.destroy()


class ScreenshotMappingDialog(tk.Toplevel):
    def __init__(
        self,
        parent: tk.Tk,
        system_names: List[str],
        screenshots: List[Path],
        initial_mapping: Dict[str, Optional[Path]],
    ):
        super().__init__(parent)
        self.title("Review screenshot mapping")
        self.resizable(True, True)
        self.result: Optional[Dict[str, Optional[Path]]] = None
        self.screenshot_by_label = {self._label(path): path for path in screenshots}
        self.vars_by_system: Dict[str, tk.StringVar] = {}

        self.transient(parent)
        self.grab_set()
        self.protocol("WM_DELETE_WINDOW", self.cancel)

        outer = ttk.Frame(self, padding=18)
        outer.grid(row=0, column=0, sticky="nsew")
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(2, weight=1)

        ttk.Label(outer, text="Review screenshot matches", style="Title.TLabel").grid(row=0, column=0, sticky="w")
        ttk.Label(
            outer,
            text=(
                "Auto-matching is intentionally conservative. Confirm each worksheet's screenshot, "
                "or leave it as (none) to skip screenshot checking for that system."
            ),
            wraplength=820,
        ).grid(row=1, column=0, sticky="we", pady=(6, 14))

        canvas = tk.Canvas(outer, highlightthickness=0)
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        table = ttk.Frame(canvas)
        table.columnconfigure(1, weight=1)
        canvas_window = canvas.create_window((0, 0), window=table, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.grid(row=2, column=0, sticky="nsew")
        scrollbar.grid(row=2, column=1, sticky="ns")

        def resize_table(event: tk.Event) -> None:
            canvas.itemconfigure(canvas_window, width=event.width)

        def update_scroll(_event: tk.Event) -> None:
            canvas.configure(scrollregion=canvas.bbox("all"))

        canvas.bind("<Configure>", resize_table)
        table.bind("<Configure>", update_scroll)

        ttk.Label(table, text="Worksheet / system").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=(0, 8))
        ttk.Label(table, text="Screenshot").grid(row=0, column=1, sticky="w", pady=(0, 8))
        options = ["(none)"] + list(self.screenshot_by_label)
        for row_index, system_name in enumerate(system_names, start=1):
            ttk.Label(table, text=system_name).grid(row=row_index, column=0, sticky="w", padx=(0, 12), pady=4)
            initial_path = initial_mapping.get(system_name)
            initial_label = self._label(initial_path) if initial_path else "(none)"
            if initial_label not in options:
                initial_label = "(none)"
            var = tk.StringVar(value=initial_label)
            self.vars_by_system[system_name] = var
            combo = ttk.Combobox(table, textvariable=var, values=options, state="readonly")
            combo.grid(row=row_index, column=1, sticky="we", pady=4)

        buttons = ttk.Frame(outer)
        buttons.grid(row=3, column=0, columnspan=2, sticky="e", pady=(16, 0))
        ttk.Button(buttons, text="Cancel", command=self.cancel).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(buttons, text="Use these mappings", command=self.accept, style="Accent.TButton").grid(row=0, column=1)

        self.bind("<Escape>", lambda _event: self.cancel())
        self.update_idletasks()
        self.minsize(760, min(680, max(420, self.winfo_height())))

    def _label(self, path: Optional[Path]) -> str:
        if path is None:
            return "(none)"
        return Path(path).name

    def accept(self) -> None:
        mapping: Dict[str, Optional[Path]] = {}
        for system_name, var in self.vars_by_system.items():
            label = var.get()
            mapping[system_name] = None if label == "(none)" else self.screenshot_by_label.get(label)
        self.result = mapping
        self.destroy()

    def cancel(self) -> None:
        self.result = None
        self.destroy()


class AuditComparatorApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("ISMS Admin Account Audit Comparator")
        self.geometry("860x560")
        self.minsize(760, 510)
        self.previous_path = tk.StringVar()
        self.current_path = tk.StringVar()
        self.screenshot_enabled = tk.BooleanVar(value=False)
        self.screenshot_path = tk.StringVar()
        self.status = tk.StringVar(value="Select the previous and current quarter Excel files.")
        self.progress = tk.IntVar(value=0)
        self._configure_style()
        self._build_ui()

    def _configure_style(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure(".", font=("Segoe UI", 10))
        style.configure("Title.TLabel", font=("Segoe UI", 16, "bold"))
        style.configure("Subtitle.TLabel", foreground="#4b5563")
        style.configure("Accent.TButton", padding=(14, 8))
        style.configure("Status.TLabel", foreground="#374151")

    def _build_ui(self) -> None:
        root = ttk.Frame(self, padding=22)
        root.grid(row=0, column=0, sticky="nsew")
        self.columnconfigure(0, weight=1)
        self.rowconfigure(0, weight=1)
        root.columnconfigure(1, weight=1)

        ttk.Label(root, text="ISMS Admin Account Audit Comparator", style="Title.TLabel").grid(
            row=0, column=0, columnspan=3, sticky="w"
        )
        ttk.Label(
            root,
            text="Compare quarterly admin account workbooks offline and generate a self-contained HTML report.",
            style="Subtitle.TLabel",
        ).grid(row=1, column=0, columnspan=3, sticky="w", pady=(6, 22))

        self._file_row(root, 2, "Previous quarter file", self.previous_path, self.choose_previous)
        self._file_row(root, 3, "Current quarter file", self.current_path, self.choose_current)

        screenshot_box = ttk.LabelFrame(root, text="Optional screenshot evidence check", padding=12)
        screenshot_box.grid(row=4, column=0, columnspan=3, sticky="we", pady=(14, 0))
        screenshot_box.columnconfigure(1, weight=1)
        ttk.Checkbutton(
            screenshot_box,
            text="Enable screenshot check",
            variable=self.screenshot_enabled,
            command=self.toggle_screenshot_controls,
        ).grid(row=0, column=0, columnspan=4, sticky="w")
        ttk.Label(screenshot_box, text="Screenshot file or folder").grid(row=1, column=0, sticky="w", pady=(10, 0))
        self.screenshot_entry = ttk.Entry(screenshot_box, textvariable=self.screenshot_path)
        self.screenshot_entry.grid(row=1, column=1, sticky="we", padx=10, pady=(10, 0))
        self.screenshot_file_button = ttk.Button(screenshot_box, text="File...", command=self.choose_screenshot_file)
        self.screenshot_file_button.grid(row=1, column=2, sticky="e", pady=(10, 0))
        self.screenshot_folder_button = ttk.Button(
            screenshot_box, text="Folder...", command=self.choose_screenshot_folder
        )
        self.screenshot_folder_button.grid(row=1, column=3, sticky="e", padx=(6, 0), pady=(10, 0))
        ttk.Label(
            screenshot_box,
            text="Folder mode auto-matches image filenames to worksheet names. OCR is local and requires Tesseract.",
            style="Subtitle.TLabel",
            wraplength=760,
        ).grid(row=2, column=0, columnspan=4, sticky="we", pady=(8, 0))
        self.toggle_screenshot_controls()

        actions = ttk.Frame(root)
        actions.grid(row=5, column=0, columnspan=3, sticky="we", pady=(24, 12))
        ttk.Button(actions, text="Run comparison and save report", command=self.run_comparison, style="Accent.TButton").grid(
            row=0, column=0, sticky="w"
        )
        ttk.Button(actions, text="Clear selections", command=self.clear_selection).grid(row=0, column=1, padx=(10, 0))

        ttk.Progressbar(root, variable=self.progress, maximum=100).grid(
            row=6, column=0, columnspan=3, sticky="we", pady=(12, 8)
        )
        ttk.Label(root, textvariable=self.status, style="Status.TLabel", wraplength=760).grid(
            row=7, column=0, columnspan=3, sticky="we"
        )

        note = (
            "Runs locally only. No network calls, cloud services, or external APIs are used. "
            "For legacy .xls files, install the optional xlrd dependency from requirements.txt."
        )
        ttk.Label(root, text=note, style="Subtitle.TLabel", wraplength=760).grid(
            row=8, column=0, columnspan=3, sticky="we", pady=(28, 0)
        )

    def _file_row(self, parent: ttk.Frame, row: int, label: str, variable: tk.StringVar, command: Any) -> None:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", pady=7)
        entry = ttk.Entry(parent, textvariable=variable)
        entry.grid(row=row, column=1, sticky="we", padx=10, pady=7)
        ttk.Button(parent, text="Browse...", command=command).grid(row=row, column=2, sticky="e", pady=7)

    def choose_previous(self) -> None:
        self._choose_file(self.previous_path, "Select previous quarter workbook")

    def choose_current(self) -> None:
        self._choose_file(self.current_path, "Select current quarter workbook")

    def choose_screenshot_file(self) -> None:
        filename = filedialog.askopenfilename(
            title="Select screenshot image",
            filetypes=[
                ("Screenshot images", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.webp"),
                ("All files", "*.*"),
            ],
        )
        if filename:
            self.screenshot_path.set(filename)
            self.screenshot_enabled.set(True)
            self.toggle_screenshot_controls()

    def choose_screenshot_folder(self) -> None:
        folder = filedialog.askdirectory(title="Select screenshot folder")
        if folder:
            self.screenshot_path.set(folder)
            self.screenshot_enabled.set(True)
            self.toggle_screenshot_controls()

    def toggle_screenshot_controls(self) -> None:
        state = "normal" if self.screenshot_enabled.get() else "disabled"
        for widget in (
            getattr(self, "screenshot_entry", None),
            getattr(self, "screenshot_file_button", None),
            getattr(self, "screenshot_folder_button", None),
        ):
            if widget is not None:
                widget.configure(state=state)

    def _choose_file(self, target: tk.StringVar, title: str) -> None:
        filename = filedialog.askopenfilename(
            title=title,
            filetypes=[("Excel workbooks", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if filename:
            target.set(filename)
            self.status.set("Ready to compare.")

    def clear_selection(self) -> None:
        self.previous_path.set("")
        self.current_path.set("")
        self.screenshot_enabled.set(False)
        self.screenshot_path.set("")
        self.toggle_screenshot_controls()
        self.progress.set(0)
        self.status.set("Selections cleared.")

    def set_status(self, text: str, progress: int) -> None:
        self.status.set(text)
        self.progress.set(progress)
        self.update_idletasks()

    def run_comparison(self) -> None:
        previous = Path(self.previous_path.get())
        current = Path(self.current_path.get())
        if not previous.exists() or not current.exists():
            messagebox.showwarning("Files required", "Please select both the previous and current quarter Excel files.")
            return
        screenshot_source: Optional[Path] = None
        if self.screenshot_enabled.get():
            screenshot_source = Path(self.screenshot_path.get())
            if not screenshot_source.exists():
                messagebox.showwarning(
                    "Screenshot source required",
                    "Screenshot check is enabled. Please select a screenshot image or folder.",
                )
                return
        report_path = filedialog.asksaveasfilename(
            title="Save HTML report",
            defaultextension=".html",
            initialfile=f"ISMS_Admin_Audit_Comparison_{_dt.datetime.now().strftime('%Y%m%d_%H%M')}.html",
            filetypes=[("HTML report", "*.html"), ("All files", "*.*")],
        )
        if not report_path:
            return

        try:
            self.set_status("Reading previous quarter workbook...", 10)
            previous_tables = load_workbook_tables(previous)
            self.set_status("Reading current quarter workbook...", 25)
            current_tables = load_workbook_tables(current)
            self.set_status("Checking column mappings...", 40)
            self.resolve_mappings(previous_tables, current_tables)
            self.set_status("Comparing systems and admin accounts...", 65)
            results = compare_workbooks(previous_tables, current_tables)
            if screenshot_source:
                screenshots = collect_screenshot_files(screenshot_source)
                if not screenshots:
                    messagebox.showwarning(
                        "No screenshots found",
                        "Screenshot check is enabled, but no supported image files were found.",
                    )
                    return
                initial_mapping = build_initial_screenshot_mapping(results, screenshots)
                dialog = ScreenshotMappingDialog(
                    self,
                    [system.name for system in results],
                    screenshots,
                    initial_mapping,
                )
                self.wait_window(dialog)
                if dialog.result is None:
                    self.status.set("Screenshot mapping cancelled.")
                    return
                self.set_status("Running local OCR screenshot checks...", 76)
                add_screenshot_checks(results, current_tables, screenshot_source, dialog.result)
            self.set_status("Building HTML report...", 82)
            report = build_html_report(previous, current, results)
            Path(report_path).write_text(report, encoding="utf-8")
            self.set_status("Report generated. File selections were cleared to prevent stale reuse.", 100)
            self.previous_path.set("")
            self.current_path.set("")
            self.screenshot_enabled.set(False)
            self.screenshot_path.set("")
            self.toggle_screenshot_controls()
            if messagebox.askyesno("Report generated", "The report was generated successfully. Open it now?"):
                webbrowser.open(Path(report_path).resolve().as_uri())
        except Exception as exc:
            self.progress.set(0)
            self.status.set("Comparison failed. See the message for details.")
            details = "".join(traceback.format_exception_only(type(exc), exc)).strip()
            messagebox.showerror("Comparison failed", details)

    def resolve_mappings(self, previous_tables: Dict[str, SheetTable], current_tables: Dict[str, SheetTable]) -> None:
        common = sorted(set(previous_tables).intersection(current_tables), key=str.casefold)
        for sheet_name in common:
            previous = previous_tables[sheet_name]
            current = current_tables[sheet_name]
            if previous.roles.get("user_id") and current.roles.get("user_id"):
                continue
            dialog = MappingDialog(self, sheet_name, previous, current)
            self.wait_window(dialog)
            if dialog.result is None or dialog.result[0]:
                previous.skipped = True
                current.skipped = True
                continue
            _skip, previous_roles, current_roles = dialog.result
            previous.roles.update(previous_roles)
            current.roles.update(current_roles)


def compare_workbooks(previous_tables: Dict[str, SheetTable], current_tables: Dict[str, SheetTable]) -> Dict[str, Any]:
    previous_names = set(previous_tables)
    current_names = set(current_tables)
    common_names = sorted(previous_names.intersection(current_names), key=str.casefold)
    results: List[SystemResult] = []

    for name in common_names:
        previous = previous_tables[name]
        current = current_tables[name]
        result = SystemResult(name=name)
        if previous.skipped or current.skipped:
            result.skipped = True
            result.skip_reason = "Skipped during manual column mapping."
            results.append(result)
            continue
        if not previous.roles.get("user_id") or not current.roles.get("user_id"):
            result.skipped = True
            result.skip_reason = "User ID column could not be detected."
            results.append(result)
            continue

        prev_users = build_user_map(previous)
        curr_users = build_user_map(current)
        for key in sorted(curr_users.keys() - prev_users.keys()):
            result.additions.append(curr_users[key])
        for key in sorted(prev_users.keys() - curr_users.keys()):
            result.removals.append(prev_users[key])
        for key in sorted(prev_users.keys() & curr_users.keys()):
            field_changes = []
            for role in COMPARABLE_ROLES:
                old = prev_users[key].values.get(role, "")
                new = curr_users[key].values.get(role, "")
                if normalize_compare_value(old) != normalize_compare_value(new):
                    field_changes.append(
                        {
                            "field": ROLE_LABELS[role],
                            "previous": old,
                            "current": new,
                        }
                    )
            if field_changes:
                result.changes.append(
                    {
                        "user_id": curr_users[key].display_id or prev_users[key].display_id,
                        "fields": field_changes,
                    }
                )
        result.clean = not result.additions and not result.removals and not result.changes
        results.append(result)

    current_only = sorted(current_names - previous_names, key=str.casefold)
    previous_only = sorted(previous_names - current_names, key=str.casefold)
    return {
        "systems": results,
        "current_only": current_only,
        "previous_only": previous_only,
        "removal_summary": cross_system_removals(results),
    }


def normalize_compare_value(value: Any) -> str:
    return re.sub(r"\s+", " ", cell_to_text(value)).strip().lower()


def cross_system_removals(results: List[SystemResult]) -> List[Dict[str, Any]]:
    by_user: Dict[str, Dict[str, Any]] = {}
    for result in results:
        for record in result.removals:
            entry = by_user.setdefault(record.key, {"user_id": record.display_id, "systems": []})
            entry["systems"].append(result.name)
    return sorted(
        [entry for entry in by_user.values() if len(entry["systems"]) > 1],
        key=lambda item: (-len(item["systems"]), item["user_id"].lower()),
    )


def add_screenshot_checks(
    data: Dict[str, Any],
    current_tables: Dict[str, SheetTable],
    source: Path,
    screenshot_mapping: Optional[Dict[str, Optional[Path]]] = None,
) -> None:
    systems: List[SystemResult] = data["systems"]
    if not systems:
        return

    if Image is None or ImageOps is None or ImageEnhance is None or pytesseract is None:
        message = "Screenshot OCR skipped. Install Pillow and pytesseract, then install the Tesseract OCR engine."
        for system in systems:
            system.screenshot_check = ScreenshotCheckResult(system_name=system.name, status="unavailable", message=message)
        return

    try:
        pytesseract.get_tesseract_version()
    except Exception:
        message = (
            "Screenshot OCR skipped. The Tesseract OCR engine was not found. "
            "Install Tesseract and make sure tesseract.exe is available on PATH."
        )
        if shutil.which("tesseract") is None:
            message += " The current PATH does not include tesseract.exe."
        for system in systems:
            system.screenshot_check = ScreenshotCheckResult(system_name=system.name, status="unavailable", message=message)
        return

    screenshots = collect_screenshot_files(source)
    if screenshot_mapping is None:
        screenshot_mapping = build_initial_screenshot_mapping(data, screenshots)
    ocr_cache: Dict[Path, str] = {}
    for system in systems:
        table = current_tables.get(system.name)
        if system.skipped or table is None or not table.roles.get("user_id"):
            system.screenshot_check = ScreenshotCheckResult(
                system_name=system.name,
                status="skipped",
                message="Screenshot check skipped because this worksheet has no usable User ID mapping.",
            )
            continue
        users = list(build_user_map(table).values())
        raw_match = screenshot_mapping.get(system.name)
        match = Path(raw_match) if raw_match else None
        if match is None:
            system.screenshot_check = ScreenshotCheckResult(
                system_name=system.name,
                status="no_screenshot",
                message="No screenshot was mapped to this worksheet.",
            )
            continue
        try:
            if match not in ocr_cache:
                ocr_cache[match] = ocr_image(match)
            text = ocr_cache[match]
            user_checks = []
            for user in users:
                present, evidence, match_type = user_id_seen_in_text(user.display_id, text)
                user_checks.append(
                    ScreenshotUserCheck(
                        user_id=user.display_id,
                        present=present,
                        evidence=evidence,
                        match_type=match_type,
                    )
                )
            missing = sum(1 for check in user_checks if not check.present)
            system.screenshot_check = ScreenshotCheckResult(
                system_name=system.name,
                screenshot_path=str(match),
                status="checked",
                message=f"{len(user_checks) - missing} of {len(user_checks)} current workbook users were found in the screenshot OCR text.",
                users=user_checks,
                ocr_preview=preview_ocr_text(text),
            )
        except Exception as exc:
            system.screenshot_check = ScreenshotCheckResult(
                system_name=system.name,
                screenshot_path=str(match),
                status="ocr_error",
                message=f"OCR failed for this screenshot: {exc}",
            )


def build_initial_screenshot_mapping(data: Dict[str, Any], screenshots: List[Path]) -> Dict[str, Optional[Path]]:
    systems: List[SystemResult] = data["systems"]
    system_count = len(systems)
    return {
        system.name: match_screenshot_to_system(system.name, screenshots, system_count)
        for system in systems
    }


def collect_screenshot_files(source: Path) -> List[Path]:
    source = Path(source)
    if source.is_file():
        return [source] if source.suffix.lower() in SCREENSHOT_EXTENSIONS else []
    if not source.is_dir():
        return []
    return sorted(
        [path for path in source.iterdir() if path.is_file() and path.suffix.lower() in SCREENSHOT_EXTENSIONS],
        key=lambda path: path.name.casefold(),
    )


def normalize_match_name(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def match_name_tokens(value: str) -> List[str]:
    return [token for token in re.split(r"[^a-z0-9]+", value.lower()) if len(token) >= 3]


def match_screenshot_to_system(system_name: str, screenshots: List[Path], system_count: int) -> Optional[Path]:
    screenshots = [Path(path) for path in screenshots]
    if not screenshots:
        return None
    target = normalize_match_name(system_name)
    if not target:
        return None
    by_exact = {normalize_match_name(path.stem): path for path in screenshots}
    if target in by_exact:
        return by_exact[target]
    target_tokens = set(match_name_tokens(system_name))
    candidates: List[Path] = []
    for path in screenshots:
        candidate = normalize_match_name(path.stem)
        if len(target) >= 5 and len(candidate) >= 5 and (candidate.startswith(target) or target.startswith(candidate)):
            candidates.append(path)
            continue
        candidate_tokens = set(match_name_tokens(path.stem))
        if target_tokens and target_tokens == candidate_tokens:
            candidates.append(path)
    return candidates[0] if len(candidates) == 1 else None


def ocr_image(path: Path) -> str:
    image = Image.open(path)
    try:
        image = ImageOps.exif_transpose(image)
        image = image.convert("L")
        image = ImageOps.autocontrast(image)
        width, height = image.size
        if width < 1800:
            scale = max(2, min(4, 1800 // max(width, 1)))
            image = image.resize((width * scale, height * scale))
        image = ImageEnhance.Contrast(image).enhance(1.8)
        image = ImageEnhance.Sharpness(image).enhance(1.5)
        return pytesseract.image_to_string(image, config="--psm 6")
    finally:
        image.close()


def normalize_ocr_text(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().lower()


def compact_ocr_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def excerpt_around(text: str, start: int, end: int, radius: int = 44) -> str:
    left = max(0, start - radius)
    right = min(len(text), end + radius)
    prefix = "..." if left > 0 else ""
    suffix = "..." if right < len(text) else ""
    return prefix + re.sub(r"\s+", " ", text[left:right]).strip() + suffix


def preview_ocr_text(text: str, limit: int = 2500) -> str:
    cleaned = re.sub(r"\s+", " ", text).strip()
    if len(cleaned) <= limit:
        return cleaned
    return cleaned[:limit].rstrip() + "..."


def compact_index_map(text: str) -> Tuple[str, List[int]]:
    compact_chars: List[str] = []
    positions: List[int] = []
    for index, char in enumerate(text):
        if char.isalnum():
            compact_chars.append(char.lower())
            positions.append(index)
    return "".join(compact_chars), positions


def user_id_seen_in_text(user_id: str, ocr_text: str) -> Tuple[bool, str, str]:
    user = cell_to_text(user_id).strip()
    if not user:
        return False, "", ""
    normalized_user = normalize_ocr_text(user)
    cleaned_text = re.sub(r"\s+", " ", ocr_text).strip()
    normalized_text = cleaned_text.lower()
    if len(normalized_user) >= 3:
        token_pattern = re.compile(r"(?<![a-z0-9])" + re.escape(normalized_user) + r"(?![a-z0-9])", re.IGNORECASE)
        match = token_pattern.search(normalized_text)
        if match:
            return True, excerpt_around(cleaned_text, match.start(), match.end()), "exact"
        loose_index = normalized_text.find(normalized_user)
        if loose_index >= 0:
            return True, excerpt_around(cleaned_text, loose_index, loose_index + len(normalized_user)), "text"
    compact_user = compact_ocr_text(user)
    compact_text, positions = compact_index_map(cleaned_text)
    if len(compact_user) >= 5:
        compact_index = compact_text.find(compact_user)
        if compact_index >= 0:
            start = positions[compact_index]
            end = positions[min(compact_index + len(compact_user) - 1, len(positions) - 1)] + 1
            return True, excerpt_around(cleaned_text, start, end), "compact"
    return False, "", ""


def request_present(record: UserRecord) -> bool:
    return bool(record.values.get("request_id", "").strip())


def html_escape(value: Any) -> str:
    return html.escape(cell_to_text(value), quote=True)


def pill(text: str, class_name: str) -> str:
    return f'<span class="pill {class_name}">{html_escape(text)}</span>'


def build_html_report(previous_path: Path, current_path: Path, data: Dict[str, Any]) -> str:
    generated = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    systems: List[SystemResult] = data["systems"]
    total_additions = sum(len(system.additions) for system in systems)
    total_removals = sum(len(system.removals) for system in systems)
    total_changes = sum(len(system.changes) for system in systems)
    clean_systems = sum(1 for system in systems if system.clean)
    missing_request_ids = sum(1 for system in systems for record in system.additions if not request_present(record))
    screenshot_checks = [system.screenshot_check for system in systems if system.screenshot_check is not None]
    screenshot_missing = sum(
        check.missing_count for check in screenshot_checks if check is not None and check.status == "checked"
    )
    finding_systems = sum(
        1
        for system in systems
        if not system.clean and not system.skipped and (system.additions or system.removals or system.changes)
    )

    cards = [
        ("Additions", total_additions, "additions"),
        ("Removals", total_removals, "removals"),
        ("Changed users", total_changes, "changes"),
        ("Clean systems", clean_systems, "clean"),
        ("Missing request IDs", missing_request_ids, "alert"),
        ("Systems with findings", finding_systems, "changes"),
    ]
    if screenshot_checks:
        cards.append(("Screenshot misses", screenshot_missing, "alert" if screenshot_missing else "clean"))
    summary_cards = "\n".join(
        f'<div class="stat" data-kind="{kind}"><strong>{value}</strong><span>{label}</span></div>'
        for label, value, kind in cards
    )

    system_sections = "\n".join(render_system(system) for system in systems)
    new_systems = render_system_list("Newly procured systems", data["current_only"], "Sheet only in current workbook.")
    decommissioned = render_system_list("Decommissioned systems", data["previous_only"], "Sheet only in previous workbook.")
    cross_removals = render_cross_removals(data["removal_summary"])

    payload = {
        "previous": previous_path.name,
        "current": current_path.name,
        "generated": generated,
    }

    return HTML_TEMPLATE.replace("{{SUMMARY_CARDS}}", summary_cards).replace(
        "{{SYSTEM_SECTIONS}}", system_sections
    ).replace("{{NEW_SYSTEMS}}", new_systems).replace("{{DECOMMISSIONED_SYSTEMS}}", decommissioned).replace(
        "{{CROSS_REMOVALS}}", cross_removals
    ).replace(
        "{{PAYLOAD}}", json.dumps(payload)
    )


def render_system_list(title: str, systems: List[str], note: str) -> str:
    if not systems:
        return ""
    items = "\n".join(f"<li>{html_escape(system)}</li>" for system in systems)
    return f"""
    <section class="panel notice-panel">
      <h2>{html_escape(title)}</h2>
      <p class="muted">{html_escape(note)}</p>
      <ul class="system-list">{items}</ul>
    </section>
    """


def render_cross_removals(entries: List[Dict[str, Any]]) -> str:
    if not entries:
        return """
        <section class="panel">
          <h2>Cross-system removals</h2>
          <p class="muted">No user was removed from multiple systems in this comparison.</p>
        </section>
        """
    rows = []
    for entry in entries:
        systems = ", ".join(entry["systems"])
        rows.append(
            f"<tr><td><code>{html_escape(entry['user_id'])}</code></td>"
            f"<td>{len(entry['systems'])}</td><td>{html_escape(systems)}</td></tr>"
        )
    return f"""
    <section class="panel cross-system">
      <h2>Cross-system removals</h2>
      <p class="muted">Users removed from more than one system. Coordinate follow-up once per user where possible.</p>
      <div class="table-wrap">
        <table>
          <thead><tr><th>User ID</th><th>Systems</th><th>System names</th></tr></thead>
          <tbody>{''.join(rows)}</tbody>
        </table>
      </div>
    </section>
    """


def render_system(system: SystemResult) -> str:
    if system.skipped:
        body = f'<p class="muted">{html_escape(system.skip_reason)}</p>' + render_screenshot_check(system)
        badge = pill("Skipped", "neutral")
        open_attr = " open"
    else:
        body = render_additions(system) + render_removals(system) + render_changes(system) + render_screenshot_check(system)
        screenshot_badge = render_screenshot_badge(system)
        if system.clean:
            body = '<p class="muted">No additions, removals, or tracked field changes detected.</p>' + render_screenshot_check(system)
            badge = pill("Clean", "ok") + screenshot_badge
            open_attr = " open" if screenshot_needs_attention(system) else ""
        else:
            badge = (
                pill(f"{len(system.additions)} additions", "alert")
                + pill(f"{len(system.removals)} removals", "warn")
                + pill(f"{len(system.changes)} changed", "info")
                + screenshot_badge
            )
            open_attr = " open"
    return f"""
    <details class="system"{open_attr}>
      <summary>
        <span>{html_escape(system.name)}</span>
        <span class="badges">{badge}</span>
      </summary>
      <div class="system-body">{body}</div>
    </details>
    """


def screenshot_needs_attention(system: SystemResult) -> bool:
    check = system.screenshot_check
    return bool(check and check.status != "checked" or check and check.missing_count > 0)


def render_screenshot_badge(system: SystemResult) -> str:
    check = system.screenshot_check
    if check is None:
        return ""
    if check.status == "checked":
        if check.missing_count:
            return pill(f"{check.missing_count} screenshot missing", "alert")
        return pill("Screenshot verified", "ok")
    if check.status == "no_screenshot":
        return pill("No screenshot", "warn")
    if check.status in {"unavailable", "ocr_error"}:
        return pill("OCR not checked", "warn")
    return pill("Screenshot skipped", "neutral")


def render_screenshot_check(system: SystemResult) -> str:
    check = system.screenshot_check
    if check is None:
        return ""
    source = ""
    if check.screenshot_path:
        source = f'<p class="muted screenshot-source">Screenshot: <code>{html_escape(Path(check.screenshot_path).name)}</code></p>'
    if check.status != "checked":
        return f"""
        <h3 class="subsection-label screenshot-label">Screenshot check</h3>
        <div class="screenshot-notice {html_escape(check.status)}">
          <p>{html_escape(check.message)}</p>
          {source}
        </div>
        """
    rows = []
    for user in check.users:
        status = pill("Found", "ok") if user.present else pill("Missing", "alert")
        evidence = ""
        if user.present:
            evidence = (
                f'<span class="match-type">{html_escape(user.match_type or "match")}</span> '
                f'{html_escape(user.evidence)}'
            )
        else:
            evidence = '<span class="empty-dash">No OCR text matched this User ID.</span>'
        rows.append(
            "<tr>"
            f"<td><code>{html_escape(user.user_id)}</code></td>"
            f"<td>{status}</td>"
            f"<td>{evidence}</td>"
            "</tr>"
        )
    preview = ""
    if check.ocr_preview:
        preview = f"""
        <details class="ocr-preview">
          <summary>View OCR text preview</summary>
          <pre>{html_escape(check.ocr_preview)}</pre>
        </details>
        """
    return f"""
    <h3 class="subsection-label screenshot-label">Screenshot check</h3>
    <p class="muted screenshot-summary">{html_escape(check.message)}</p>
    {source}
    <div class="table-wrap">
      <table>
        <thead><tr><th>User ID from current workbook</th><th>Seen in screenshot OCR</th><th>OCR evidence</th></tr></thead>
        <tbody>{''.join(rows)}</tbody>
      </table>
    </div>
    {preview}
    """


def render_additions(system: SystemResult) -> str:
    if not system.additions:
        return ""
    rows = []
    for record in system.additions:
        request_id = record.values.get("request_id", "")
        status = pill("Provided", "ok") if request_present(record) else pill("Missing", "alert")
        rows.append(
            "<tr>"
            f"<td><code>{html_escape(record.display_id)}</code></td>"
            f"<td>{html_escape(record.values.get('rights_type', ''))}</td>"
            f"<td>{html_escape(record.values.get('purpose', ''))}</td>"
            f"<td><code>{html_escape(request_id)}</code></td>"
            f"<td>{status}</td>"
            "</tr>"
        )
    return f"""
    <h3 class="subsection-label additions-label">Additions</h3>
    <div class="table-wrap">
      <table>
        <thead><tr><th>User ID</th><th>Admin rights</th><th>Purpose</th><th>Request ID</th><th>Request status</th></tr></thead>
        <tbody>{''.join(rows)}</tbody>
      </table>
    </div>
    """


def render_removals(system: SystemResult) -> str:
    if not system.removals:
        return ""
    rows = []
    for record in system.removals:
        rows.append(
            "<tr>"
            f"<td><code>{html_escape(record.display_id)}</code></td>"
            f"<td>{html_escape(record.values.get('rights_type', ''))}</td>"
            f"<td>{html_escape(record.values.get('purpose', ''))}</td>"
            f"<td>{pill('Justification required', 'warn')}</td>"
            "</tr>"
        )
    return f"""
    <h3 class="subsection-label removals-label">Removals</h3>
    <div class="table-wrap">
      <table>
        <thead><tr><th>User ID</th><th>Previous admin rights</th><th>Previous purpose</th><th>Status</th></tr></thead>
        <tbody>{''.join(rows)}</tbody>
      </table>
    </div>
    """


def render_changes(system: SystemResult) -> str:
    if not system.changes:
        return ""
    blocks = []
    for change in system.changes:
        rows = []
        for field in change["fields"]:
            rows.append(
                "<tr>"
                f"<td>{html_escape(field['field'])}</td>"
                f"<td class=\"diff-prev\">{html_escape(field['previous'])}</td>"
                f"<td class=\"diff-curr\">{html_escape(field['current'])}</td>"
                "</tr>"
            )
        blocks.append(
            f"""
            <div class="change-block">
              <h4 class="change-user"><code>{html_escape(change['user_id'])}</code></h4>
              <div class="table-wrap">
                <table>
                  <thead><tr><th>Field</th><th>Previous</th><th>Current</th></tr></thead>
                  <tbody>{''.join(rows)}</tbody>
                </table>
              </div>
            </div>
            """
        )
    return f"<h3 class=\"subsection-label changes-label\">Changes</h3>{''.join(blocks)}"


HTML_TEMPLATE = r"""<!doctype html>
<html lang="en" data-theme="auto">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>ISMS Admin Account Audit Comparison</title>
  <style>
    :root,
    [data-theme="light"] {
      --bg: #f7f7f7;
      --surface: #ffffff;
      --surface-raised: #ffffff;
      --surface-hover: #f2f2f2;
      --surface-inset: #ededec;
      --text-primary: #1a1a1a;
      --text-secondary: #555555;
      --text-tertiary: #777777;
      --border: #dedede;
      --border-subtle: #ebebeb;
      --accent: #5c5145;
      --accent-hover: #6e6255;
      --red: #9e3b2d;
      --red-surface: #fdf1ee;
      --red-border: #e8cac3;
      --red-text: #7a2e22;
      --amber: #7d6022;
      --amber-surface: #fdf7eb;
      --amber-border: #e2d4ac;
      --amber-text: #614a1a;
      --green: #3e7652;
      --green-surface: #f0f8f2;
      --green-border: #bdd8c6;
      --green-text: #2d5a3d;
      --blue: #3d6284;
      --blue-surface: #eff5fa;
      --blue-border: #b9cfe1;
      --blue-text: #2e4d68;
      --shadow-sm: 0 1px 2px rgba(0,0,0,0.04);
      --shadow-md: 0 2px 8px rgba(0,0,0,0.06);
      --shadow-lg: 0 4px 16px rgba(0,0,0,0.08);
      color-scheme: light;
    }
    [data-theme="dark"] {
      --bg: #1c1916;
      --surface: #262220;
      --surface-raised: #302b28;
      --surface-hover: #3a3430;
      --surface-inset: #1c1916;
      --text-primary: #ece5dc;
      --text-secondary: #a89e94;
      --text-tertiary: #8f867e;
      --border: #3e3832;
      --border-subtle: #332e2a;
      --accent: #c4ac93;
      --accent-hover: #d4bea7;
      --red: #d4745e;
      --red-surface: #2e2220;
      --red-border: #4a2f28;
      --red-text: #e8a090;
      --amber: #c8a84e;
      --amber-surface: #2a2618;
      --amber-border: #4a3f1e;
      --amber-text: #dcc26e;
      --green: #6aae7e;
      --green-surface: #1e2a22;
      --green-border: #2a4832;
      --green-text: #8ecaa0;
      --blue: #6a9cc4;
      --blue-surface: #1e2630;
      --blue-border: #2a3e52;
      --blue-text: #8ebadc;
      --shadow-sm: 0 1px 2px rgba(0,0,0,0.2);
      --shadow-md: 0 2px 8px rgba(0,0,0,0.25);
      --shadow-lg: 0 4px 16px rgba(0,0,0,0.3);
      color-scheme: dark;
    }
    @media (prefers-color-scheme: dark) {
      [data-theme="auto"] {
        --bg: #1c1916;
        --surface: #262220;
        --surface-raised: #302b28;
        --surface-hover: #3a3430;
        --surface-inset: #1c1916;
        --text-primary: #ece5dc;
        --text-secondary: #a89e94;
        --text-tertiary: #8f867e;
        --border: #3e3832;
        --border-subtle: #332e2a;
        --accent: #c4ac93;
        --accent-hover: #d4bea7;
        --red: #d4745e;
        --red-surface: #2e2220;
        --red-border: #4a2f28;
        --red-text: #e8a090;
        --amber: #c8a84e;
        --amber-surface: #2a2618;
        --amber-border: #4a3f1e;
        --amber-text: #dcc26e;
        --green: #6aae7e;
        --green-surface: #1e2a22;
        --green-border: #2a4832;
        --green-text: #8ecaa0;
        --blue: #6a9cc4;
        --blue-surface: #1e2630;
        --blue-border: #2a3e52;
        --blue-text: #8ebadc;
        --shadow-sm: 0 1px 2px rgba(0,0,0,0.2);
        --shadow-md: 0 2px 8px rgba(0,0,0,0.25);
        --shadow-lg: 0 4px 16px rgba(0,0,0,0.3);
        color-scheme: dark;
      }
    }
    :root {
      --font-body: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
      --font-mono: "SF Mono", "Cascadia Code", "Fira Code", Consolas, "Liberation Mono", monospace;
      --text-2xs: 0.6875rem;
      --text-xs: 0.8125rem;
      --text-sm: 0.9375rem;
      --text-base: 1.0625rem;
      --text-lg: 1.3125rem;
      --text-xl: 1.75rem;
      --text-2xl: 2.2rem;
      --space-1: 4px;
      --space-2: 8px;
      --space-3: 12px;
      --space-4: 16px;
      --space-5: 24px;
      --space-6: 32px;
      --space-8: 48px;
      --space-9: 64px;
      --radius-sm: 6px;
      --radius-md: 10px;
      --radius-lg: 14px;
      --radius-pill: 100px;
      --duration-fast: 120ms;
      --duration-normal: 200ms;
      --easing-out: cubic-bezier(0.16, 1, 0.3, 1);
    }
    *, *::before, *::after { box-sizing: border-box; }
    html { font-size: 16px; -webkit-text-size-adjust: 100%; }
    body {
      margin: 0;
      background: var(--bg);
      color: var(--text-primary);
      font-family: var(--font-body);
      font-size: var(--text-base);
      line-height: 1.55;
      -webkit-font-smoothing: antialiased;
      -moz-osx-font-smoothing: grayscale;
      transition: background var(--duration-normal) var(--easing-out), color var(--duration-normal) var(--easing-out);
    }
    a, button, summary { touch-action: manipulation; }
    button:focus-visible, summary:focus-visible {
      outline: 2px solid var(--accent);
      outline-offset: 2px;
    }
    .page {
      width: min(1060px, calc(100% - 32px));
      margin: 0 auto;
      padding: var(--space-8) 0 var(--space-6);
    }
    .theme-toggle-wrap {
      position: fixed;
      top: var(--space-4);
      right: var(--space-4);
      z-index: 100;
    }
    .theme-toggle {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      gap: var(--space-2);
      min-height: 38px;
      min-width: 38px;
      padding: var(--space-1) var(--space-3);
      border: 1px solid var(--border);
      border-radius: var(--radius-pill);
      background: var(--surface-raised);
      color: var(--text-secondary);
      box-shadow: var(--shadow-md);
      cursor: pointer;
      font: inherit;
      font-size: var(--text-2xs);
      font-weight: 700;
      transition: background var(--duration-fast) var(--easing-out), box-shadow var(--duration-fast) var(--easing-out);
    }
    .theme-toggle:hover { background: var(--surface-hover); box-shadow: var(--shadow-lg); }
    .theme-toggle svg { width: 18px; height: 18px; flex: 0 0 auto; }
    .theme-toggle-label { white-space: nowrap; }
    .report-header {
      margin-bottom: var(--space-8);
      padding-right: var(--space-9);
    }
    .eyebrow {
      margin: 0 0 var(--space-2);
      color: var(--accent);
      font-size: var(--text-2xs);
      font-weight: 700;
      letter-spacing: .1em;
      text-transform: uppercase;
    }
    h1 {
      margin: 0;
      color: var(--text-primary);
      font-size: var(--text-2xl);
      line-height: 1.25;
      letter-spacing: 0;
      font-weight: 750;
    }
    .meta {
      display: flex;
      flex-wrap: wrap;
      gap: var(--space-2) var(--space-6);
      margin-top: var(--space-5);
      color: var(--text-tertiary);
      font-size: var(--text-xs);
    }
    .meta span {
      display: inline-flex;
      gap: var(--space-2);
      align-items: baseline;
    }
    .meta strong {
      color: var(--text-secondary);
      font-weight: 700;
    }
    .stats {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(155px, 1fr));
      gap: var(--space-3);
      margin-bottom: var(--space-8);
    }
    .stat, .panel, details.system {
      background: var(--surface-raised);
      border: 1px solid var(--border);
      border-radius: var(--radius-md);
      box-shadow: var(--shadow-sm);
    }
    .stat {
      position: relative;
      min-height: 88px;
      padding: var(--space-4) var(--space-5);
      overflow: hidden;
    }
    .stat::after {
      content: "";
      position: absolute;
      left: var(--space-4);
      right: var(--space-4);
      bottom: 0;
      height: 2px;
      border-radius: 2px;
      background: var(--accent);
    }
    .stat[data-kind="additions"]::after, .stat[data-kind="alert"]::after { background: var(--red); }
    .stat[data-kind="removals"]::after { background: var(--amber); }
    .stat[data-kind="changes"]::after { background: var(--blue); }
    .stat[data-kind="clean"]::after { background: var(--green); }
    .stat strong {
      display: block;
      margin-bottom: var(--space-1);
      font-size: var(--text-xl);
      line-height: 1;
      font-weight: 800;
    }
    .stat[data-kind="additions"] strong, .stat[data-kind="alert"] strong { color: var(--red); }
    .stat[data-kind="removals"] strong { color: var(--amber); }
    .stat[data-kind="changes"] strong { color: var(--blue); }
    .stat[data-kind="clean"] strong { color: var(--green); }
    .stat span {
      color: var(--text-tertiary);
      font-size: var(--text-xs);
      font-weight: 600;
    }
    .panel {
      margin-bottom: var(--space-5);
      padding: var(--space-5);
      border-radius: var(--radius-lg);
    }
    h2 {
      margin: 0 0 var(--space-2);
      color: var(--text-primary);
      font-size: var(--text-lg);
      line-height: 1.25;
    }
    .muted {
      margin: 0;
      color: var(--text-secondary);
      font-size: var(--text-sm);
    }
    .system-list {
      margin: var(--space-3) 0 0;
      padding-left: var(--space-5);
      color: var(--text-primary);
      font-size: var(--text-sm);
    }
    .notice-panel { background: var(--surface-raised); }
    .cross-system {
      background: var(--amber-surface);
      border-color: var(--amber-border);
    }
    .cross-system h2 { color: var(--amber-text); }
    .systems-heading {
      margin: 0 0 var(--space-5);
      padding-bottom: var(--space-3);
      border-bottom: 2px solid var(--accent);
      font-size: var(--text-lg);
      font-weight: 750;
    }
    details.system {
      margin-bottom: var(--space-3);
      border-radius: var(--radius-lg);
      overflow: clip;
      transition: box-shadow var(--duration-normal) var(--easing-out), border-color var(--duration-fast) var(--easing-out);
    }
    details.system:hover { box-shadow: var(--shadow-md); }
    summary {
      min-height: 54px;
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: var(--space-3);
      padding: var(--space-4) var(--space-5);
      cursor: pointer;
      color: var(--text-primary);
      font-size: var(--text-sm);
      font-weight: 700;
      transition: background var(--duration-fast) var(--easing-out);
    }
    summary:hover { background: var(--surface-hover); }
    summary::marker { color: var(--text-tertiary); }
    .badges {
      display: flex;
      flex-wrap: wrap;
      justify-content: flex-end;
      gap: var(--space-2);
      font-weight: 700;
    }
    .pill {
      display: inline-flex;
      align-items: center;
      min-height: 24px;
      padding: 2px var(--space-2);
      border-radius: var(--radius-pill);
      border: 1px solid;
      font-size: var(--text-2xs);
      font-weight: 700;
      letter-spacing: .02em;
      white-space: nowrap;
    }
    .alert { background: var(--red-surface); color: var(--red-text); border-color: var(--red-border); }
    .warn { background: var(--amber-surface); color: var(--amber-text); border-color: var(--amber-border); }
    .info { background: var(--blue-surface); color: var(--blue-text); border-color: var(--blue-border); }
    .ok { background: var(--green-surface); color: var(--green-text); border-color: var(--green-border); }
    .neutral { background: var(--surface-inset); color: var(--text-tertiary); border-color: var(--border); }
    .system-body {
      border-top: 1px solid var(--border-subtle);
      padding: var(--space-5);
    }
    .subsection-label {
      margin: 0 0 var(--space-3);
      padding-bottom: var(--space-2);
      border-bottom: 1px solid;
      font-size: var(--text-2xs);
      font-weight: 800;
      letter-spacing: .07em;
      text-transform: uppercase;
    }
    .additions-label { color: var(--red-text); border-color: var(--red-border); }
    .removals-label { color: var(--amber-text); border-color: var(--amber-border); margin-top: var(--space-6); }
    .changes-label { color: var(--blue-text); border-color: var(--blue-border); margin-top: var(--space-6); }
    .screenshot-label { color: var(--green-text); border-color: var(--green-border); margin-top: var(--space-6); }
    .screenshot-summary, .screenshot-source { margin-bottom: var(--space-3); }
    .screenshot-notice {
      padding: var(--space-4) var(--space-5);
      border: 1px solid var(--border);
      border-radius: var(--radius-md);
      background: var(--surface-inset);
      color: var(--text-secondary);
      font-size: var(--text-sm);
    }
    .screenshot-notice p { margin: 0; }
    .screenshot-notice.no_screenshot,
    .screenshot-notice.unavailable,
    .screenshot-notice.ocr_error {
      background: var(--amber-surface);
      border-color: var(--amber-border);
      color: var(--amber-text);
    }
    .match-type {
      display: inline-flex;
      margin-right: var(--space-1);
      padding: 1px var(--space-2);
      border-radius: var(--radius-pill);
      background: var(--surface-inset);
      color: var(--text-secondary);
      font-size: var(--text-2xs);
      font-weight: 800;
      text-transform: uppercase;
      letter-spacing: .04em;
    }
    .ocr-preview {
      margin-top: var(--space-4);
      border: 1px solid var(--border);
      border-radius: var(--radius-md);
      background: var(--surface-inset);
      overflow: clip;
    }
    .ocr-preview summary {
      min-height: 42px;
      padding: var(--space-3) var(--space-4);
      font-size: var(--text-xs);
    }
    .ocr-preview pre {
      margin: 0;
      padding: var(--space-4);
      border-top: 1px solid var(--border-subtle);
      white-space: pre-wrap;
      color: var(--text-secondary);
      font: 600 var(--text-2xs) / 1.55 var(--font-mono);
    }
    .empty-dash { color: var(--text-tertiary); }
    .table-wrap {
      width: 100%;
      overflow-x: auto;
      -webkit-overflow-scrolling: touch;
    }
    table {
      width: 100%;
      border-collapse: collapse;
      min-width: 680px;
      background: transparent;
      font-size: var(--text-xs);
      line-height: 1.55;
    }
    th, td {
      padding: var(--space-2) var(--space-3);
      text-align: left;
      vertical-align: top;
      border-bottom: 1px solid var(--border-subtle);
    }
    th {
      color: var(--text-tertiary);
      font-size: var(--text-2xs);
      text-transform: uppercase;
      letter-spacing: .04em;
      font-weight: 800;
      border-bottom: 2px solid var(--border);
      white-space: nowrap;
    }
    tbody tr { transition: background var(--duration-fast) var(--easing-out); }
    tbody tr:hover { background: var(--surface-hover); }
    tr:last-child td { border-bottom: 0; }
    code {
      font-family: var(--font-mono);
      font-size: var(--text-2xs);
      font-weight: 700;
      overflow-wrap: anywhere;
    }
    .change-block {
      margin-bottom: var(--space-4);
    }
    .change-user {
      margin: 0 0 var(--space-2);
      color: var(--text-primary);
      font-size: var(--text-xs);
    }
    .diff-prev {
      color: var(--red);
      text-decoration: line-through;
      opacity: .78;
    }
    .diff-curr {
      color: var(--green);
      font-weight: 700;
    }
    .report-footer {
      margin-top: var(--space-9);
      padding-top: var(--space-5);
      border-top: 1px solid var(--border);
      color: var(--text-tertiary);
      text-align: center;
      font-size: var(--text-2xs);
      letter-spacing: .02em;
    }
    @media (max-width: 920px) {
      .stats { grid-template-columns: repeat(2, minmax(0, 1fr)); }
    }
    @media (max-width: 560px) {
      .page { width: min(100% - 24px, 1060px); padding-top: var(--space-6); }
      .report-header { padding-right: var(--space-8); }
      h1 { font-size: var(--text-xl); }
      .meta { flex-direction: column; gap: var(--space-1); }
      .stats { grid-template-columns: 1fr; }
      summary { align-items: flex-start; flex-direction: column; }
      .badges { justify-content: flex-start; }
      .theme-toggle-label { display: none; }
    }
    @media (prefers-reduced-motion: reduce) {
      *, *::before, *::after {
        animation-duration: 0.01ms !important;
        animation-iteration-count: 1 !important;
        scroll-behavior: auto !important;
        transition-duration: 0.01ms !important;
      }
    }
    @media print {
      :root, [data-theme="light"], [data-theme="dark"], [data-theme="auto"] {
        --bg: #ffffff;
        --surface: #ffffff;
        --surface-raised: #ffffff;
        --surface-hover: #ffffff;
        --shadow-sm: none;
        --shadow-md: none;
        --shadow-lg: none;
        color-scheme: light;
      }
      html, body { background: #fff; color: #000; }
      body { font-size: 11pt; }
      .page { width: 100%; padding: 0; }
      .theme-toggle-wrap { display: none; }
      .stat, .panel, details.system { box-shadow: none; break-inside: avoid; border: 1px solid #ccc; }
      details.system { display: block; }
      details.system > summary { list-style: none; }
      details.system:not([open]) > .system-body { display: block; }
      .table-wrap { overflow: visible; }
      table { min-width: 0; font-size: 10pt; }
      .report-header { padding-right: 0; }
    }
  </style>
</head>
<body>
  <div class="theme-toggle-wrap">
    <button class="theme-toggle" id="themeToggle" type="button" aria-label="Switch report theme">
      <svg id="themeIcon" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"></svg>
      <span class="theme-toggle-label" id="themeLabel">Auto</span>
    </button>
  </div>
  <main class="page">
    <header class="report-header">
      <p class="eyebrow">ISMS Quarterly Audit</p>
      <h1>Admin Account Comparison</h1>
      <p class="meta" id="meta"></p>
    </header>
    <section class="stats" aria-label="Summary statistics">
      {{SUMMARY_CARDS}}
    </section>
    {{NEW_SYSTEMS}}
    {{DECOMMISSIONED_SYSTEMS}}
    {{CROSS_REMOVALS}}
    <section aria-label="Per-system comparison">
      <h2 class="systems-heading">Systems</h2>
      {{SYSTEM_SECTIONS}}
    </section>
    <footer class="report-footer">
      ISMS Admin Account Audit Comparator &middot; Processed locally
    </footer>
  </main>
  <script>
    const payload = {{PAYLOAD}};
    const themes = ["light", "dark", "auto"];
    const labels = { light: "Light", dark: "Dark", auto: "Auto" };
    const icons = {
      light: '<circle cx="12" cy="12" r="5"/><path d="M12 1v2M12 21v2M4.22 4.22l1.42 1.42M18.36 18.36l1.42 1.42M1 12h2M21 12h2M4.22 19.78l1.42-1.42M18.36 5.64l1.42-1.42"/>',
      dark: '<path d="M21 12.79A9 9 0 1 1 11.21 3 7 7 0 0 0 21 12.79z"/>',
      auto: '<circle cx="12" cy="12" r="9"/><path d="M12 3a9 9 0 0 1 0 18" fill="currentColor" opacity="0.5"/>'
    };
    function storedTheme() {
      try { return localStorage.getItem("isms-report-theme"); } catch (error) { return null; }
    }
    function saveTheme(theme) {
      try { localStorage.setItem("isms-report-theme", theme); } catch (error) {}
    }
    function applyTheme(theme) {
      document.documentElement.dataset.theme = theme;
      document.getElementById("themeLabel").textContent = labels[theme];
      document.getElementById("themeIcon").innerHTML = icons[theme];
    }
    document.getElementById("meta").innerHTML =
      `<span><strong>Generated</strong>${payload.generated}</span>` +
      `<span><strong>Previous</strong>${payload.previous}</span>` +
      `<span><strong>Current</strong>${payload.current}</span>`;
    const initialTheme = themes.includes(storedTheme()) ? storedTheme() : "auto";
    applyTheme(initialTheme);
    document.getElementById("themeToggle").addEventListener("click", () => {
      const current = document.documentElement.dataset.theme || "auto";
      const next = themes[(themes.indexOf(current) + 1) % themes.length];
      applyTheme(next);
      saveTheme(next);
    });
    window.addEventListener("beforeprint", () => {
      document.querySelectorAll("details.system").forEach(detail => detail.setAttribute("open", ""));
    });
  </script>
</body>
</html>
"""


def main() -> int:
    app = AuditComparatorApp()
    app.mainloop()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
